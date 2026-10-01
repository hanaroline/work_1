# -*- coding: utf-8 -*-
"""퇴직연금 DC 제안서 — 편집용 PPTX (16:9, 표지+30/50/70%).
데이터는 gen_html.py와 동일. PowerPoint(Windows)에서 맑은 고딕으로 렌더되도록 EA 폰트 지정."""
from pptx import Presentation
from pptx.util import Inches as IN, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn
import copy

OUT="/home/user/work_1/dc-proposal/퇴직연금DC_제안서_위험자산비중별_202607.pptx"
FONT="맑은 고딕"

# ---------- 데이터 ----------
ETF={
 "반도체":dict(name="TIGER 미국필라델피아반도체나스닥",grade="매우높은위험",gs="매우높음",
   strat="미국 상장 반도체 시총 상위 30종목 · AI 인프라(전력·데이터센터) 실적 가시성 최선호"),
 "AI":dict(name="TIGER 글로벌AI액티브",grade="높은위험",gs="높음",
   strat="AI 성장 수혜 글로벌 기업에 액티브 투자 · 공급망 전반 확산 수혜"),
 "반도체TOP10":dict(name="TIGER 반도체TOP10",grade="매우높은위험",gs="매우높음",
   strat="국내 반도체 시총 상위 10종목 · 메모리 부족·CAPEX 확산의 한국 반도체 집중 수혜"),
 "차이나":dict(name="TIGER 차이나반도체FACTSET",grade="매우높은위험",gs="매우높음",
   strat="중국·홍콩 상장 반도체 시총 상위 25종목 · H200 규제에 따른 반도체 자급·AI 인프라 수혜"),
}
RETURNS={
 "반도체":["-7.72","+36.51","+62.40","+133.88"],
 "AI":["-12.27","+23.93","+45.90","+94.51"],
 "반도체TOP10":["-32.23","+3.82","+54.79","+198.59"],
 "차이나":["+19.25","+71.73","+71.26","+187.44"],
}
PORTS=[
 dict(tag="30%",risk=30,profile="안정추구형",
   desc="원금 안정성을 최우선으로, 위험자산은 최소한으로 가져가되 핵심 성장 테마에만 집중",
   items=[("반도체",12),("AI",9),("반도체TOP10",9)]),
 dict(tag="50%",risk=50,profile="위험중립형",
   desc="수익과 안정의 균형. 안전자산 절반으로 변동성을 제어하며 성장 테마를 폭넓게 편입",
   items=[("반도체",20),("AI",15),("반도체TOP10",11),("차이나",4)]),
 dict(tag="70%",risk=70,profile="수익추구형",
   desc="퇴직연금 위험자산 최대 한도(70%)를 활용해 반도체·AI 성장 테마의 수익 기회를 적극 추구",
   items=[("반도체",28),("AI",22),("반도체TOP10",14),("차이나",6)]),
]
DEPOSIT=dict(name="한국증권금융 정기예금",grade="매우낮은위험",gs="매우낮음",
   strat="원리금보장 · 예금자보호(1인 1억원) · AAA 최고 신용등급 · 위험자산 변동성에 대한 안전판(확정금리)",rate="3.70")

NAVY=RGBColor(0x14,0x2A,0x4C); ORANGE=RGBColor(0xF2,0x6A,0x1B); GOLD=RGBColor(0xF2,0x9F,0x05)
SAFE=RGBColor(0x2E,0x6D,0xB4); WHITE=RGBColor(0xFF,0xFF,0xFF); DGRAY=RGBColor(0x3a,0x44,0x50)
RUP=RGBColor(0xC0,0x39,0x2B); RDN=RGBColor(0x1F,0x6F,0xB2); LGRAY=RGBColor(0x9a,0xa4,0xb2)
HEADBG=RGBColor(0xEE,0xF2,0xF8); DEPBG=RGBColor(0xF3,0xF8,0xF4)
GRADE_COLOR={"매우높은위험":RGBColor(0xB2,0x3A,0x2E),"높은위험":RGBColor(0xD0,0x6A,0x1B),
             "매우낮은위험":RGBColor(0x2E,0x8B,0x6B)}

def ret_color(v):
    v=v.strip()
    return RUP if v.startswith("+") else (RDN if v.startswith("-") else DGRAY)

def set_ea(run):
    """맑은 고딕을 동아시아 글꼴로도 지정."""
    rPr=run._r.get_or_add_rPr()
    for tag in ("a:latin","a:ea","a:cs"):
        e=rPr.find(qn(tag))
        if e is None:
            e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set("typeface",FONT)

def style_run(r,size,bold=False,color=None):
    r.font.size=Pt(size); r.font.bold=bold; r.font.name=FONT
    if color is not None: r.font.color.rgb=color
    set_ea(r)

def add_text(slide,l,t,w,h,lines,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tb=slide.shapes.add_textbox(l,t,w,h); tf=tb.text_frame; tf.word_wrap=True
    tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,(txt,sz,bold,col) in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align
        r=p.add_run(); r.text=txt; style_run(r,sz,bold,col)
    return tb

def fill(shape,color,line=None):
    shape.fill.solid(); shape.fill.fore_color.rgb=color
    if line is None: shape.line.fill.background()
    else: shape.line.color.rgb=line; shape.line.width=Pt(0.75)

prs=Presentation(); prs.slide_width=IN(13.333); prs.slide_height=IN(7.5)
blank=prs.slide_layouts[6]
SW=13.333

# ============ 표지 ============
s=prs.slides.add_slide(blank)
bg=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,IN(13.333),IN(7.5)); fill(bg,NAVY)
bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,IN(0.9),IN(2.4),IN(1.3),IN(0.14)); fill(bar,ORANGE)
add_text(s,IN(0.9),IN(0.7),IN(8),IN(0.5),[("미래에셋증권 마포WM",18,True,RGBColor(0xCD,0xD8,0xEA))])
add_text(s,IN(0.9),IN(2.7),IN(11),IN(1.6),[("퇴직연금 DC 제안서",54,True,WHITE)])
add_text(s,IN(0.92),IN(4.25),IN(11),IN(0.7),[("위험자산 비중별 포트폴리오 제안",26,False,RGBColor(0xD7,0xDE,0xEA))])
# 칩
chips=[("위험자산 30%",ORANGE,WHITE),("50%",None,RGBColor(0xEA,0xF0,0xF8)),("70%",None,RGBColor(0xEA,0xF0,0xF8))]
cx=IN(0.9)
for txt,bgc,fgc in chips:
    w=IN(1.9) if "위험" in txt else IN(1.1)
    c=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,cx,IN(5.25),w,IN(0.62))
    if bgc: fill(c,bgc)
    else:
        c.fill.background(); c.line.color.rgb=RGBColor(0xAA,0xBB,0xD5); c.line.width=Pt(1)
    tf=c.text_frame; tf.word_wrap=False; p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=txt; style_run(r,15,True,fgc)
    cx=Emu(cx+w+IN(0.25))
add_text(s,IN(0.9),IN(6.6),IN(6),IN(0.4),[("미래에셋증권 마포WM",15,False,RGBColor(0xB9,0xC6,0xDC))])
add_text(s,IN(9.4),IN(6.6),IN(3),IN(0.4),[("2026년 7월",15,False,RGBColor(0xB9,0xC6,0xDC))],align=PP_ALIGN.RIGHT)

# ============ 본문 ============
HEADERS=["구분","상품명 · 투자전략","위험등급","비중\n(%)","금리\n(연,%,세전)","1개월","3개월","6개월","1년"]
COLW=[IN(0.75),IN(3.35),IN(1.15),IN(0.75),IN(1.05),IN(0.78),IN(0.78),IN(0.78),IN(0.78)]  # 표 좌측폭

def content(p):
    s=prs.slides.add_slide(blank)
    # 밴드
    band=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,IN(13.333),IN(1.25)); fill(band,NAVY)
    add_text(s,IN(0.5),IN(0.16),IN(9),IN(0.35),[("퇴직연금 DC 포트폴리오 제안",13,True,RGBColor(0xB9,0xC6,0xDC))])
    tb=s.shapes.add_textbox(IN(0.5),IN(0.5),IN(9),IN(0.65)); tf=tb.text_frame; tf.word_wrap=True
    tf.margin_top=0;tf.margin_bottom=0;tf.margin_left=0
    pp=tf.paragraphs[0]
    r=pp.add_run(); r.text=f"위험자산 비중 {p['risk']}% · "; style_run(r,24,True,WHITE)
    r=pp.add_run(); r.text=p["profile"]; style_run(r,24,True,GOLD)
    tag=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,IN(11.3),IN(0.33),IN(1.55),IN(0.6)); fill(tag,ORANGE)
    tf=tag.text_frame; pr=tf.paragraphs[0]; pr.alignment=PP_ALIGN.CENTER
    r=pr.add_run(); r.text=p["tag"]; style_run(r,28,True,WHITE)

    # ----- 표 -----
    rows=p["items"]; ndata=len(rows)+1  # +예금
    nrows=2+ndata  # 2행 헤더
    left=IN(0.5); top=IN(1.6)
    tw=sum(COLW,Emu(0))
    th=IN(0.62)+IN(0.42)*ndata  # 대략
    gt=s.shapes.add_table(nrows,9,left,top,tw,IN(0.5)).table
    # 열 폭
    for i,w in enumerate(COLW): gt.columns[i].width=w
    # 헤더 높이
    gt.rows[0].height=IN(0.42); gt.rows[1].height=IN(0.30)
    # 헤더 병합: 0..4 열은 2행 세로병합, 5~8은 상단 1칸 가로병합(기간별 수익률)
    for c in range(5):
        gt.cell(0,c).merge(gt.cell(1,c))
    gt.cell(0,5).merge(gt.cell(0,8))
    def cell_text(cell,txt,sz,bold,color,align=PP_ALIGN.CENTER,bg=None):
        cell.vertical_anchor=MSO_ANCHOR.MIDDLE
        cell.margin_left=Pt(3);cell.margin_right=Pt(3);cell.margin_top=Pt(1);cell.margin_bottom=Pt(1)
        if bg is not None: cell.fill.solid(); cell.fill.fore_color.rgb=bg
        else: cell.fill.solid(); cell.fill.fore_color.rgb=WHITE
        tf=cell.text_frame; tf.word_wrap=True
        p0=tf.paragraphs[0]; p0.alignment=align
        first=True
        for ln in txt.split("\n"):
            pp=p0 if first else tf.add_paragraph(); pp.alignment=align; first=False
            r=pp.add_run(); r.text=ln; style_run(r,sz,bold,color)
    # 헤더 채우기
    cell_text(gt.cell(0,0),"구분",11,True,NAVY,bg=HEADBG)
    cell_text(gt.cell(0,1),"상품명 · 투자전략",11,True,NAVY,bg=HEADBG)
    cell_text(gt.cell(0,2),"위험등급",11,True,NAVY,bg=HEADBG)
    cell_text(gt.cell(0,3),"비중(%)",11,True,NAVY,bg=HEADBG)
    cell_text(gt.cell(0,4),"금리\n(연,%,세전)",10,True,NAVY,bg=HEADBG)
    cell_text(gt.cell(0,5),"기간별 수익률 (%)",11,True,NAVY,bg=RGBColor(0xE7,0xED,0xF6))
    for i,lab in enumerate(["1개월","3개월","6개월","1년"]):
        cell_text(gt.cell(1,5+i),lab,10,True,NAVY,bg=RGBColor(0xE7,0xED,0xF6))
    # 데이터행
    def fillrow(ri,gub,e,w,rate,rets,dep=False):
        bg=DEPBG if dep else WHITE
        gt.rows[ri].height=IN(0.72)
        cell_text(gt.cell(ri,0),gub,12,True,NAVY,bg=bg)
        # 상품명+전략 (2문단)
        c=gt.cell(ri,1); c.vertical_anchor=MSO_ANCHOR.MIDDLE
        c.margin_left=Pt(5);c.margin_right=Pt(4);c.margin_top=Pt(2);c.margin_bottom=Pt(2)
        c.fill.solid(); c.fill.fore_color.rgb=bg
        tf=c.text_frame; tf.word_wrap=True
        p0=tf.paragraphs[0]; p0.alignment=PP_ALIGN.LEFT
        r=p0.add_run(); r.text=e["name"]; style_run(r,12,True, RGBColor(0x1D,0x5B,0x43) if dep else RGBColor(0x15,0x23,0x3b))
        p1=tf.add_paragraph(); p1.alignment=PP_ALIGN.LEFT
        r=p1.add_run(); r.text=e["strat"]; style_run(r,8.5,False,RGBColor(0x5a,0x66,0x76))
        # 위험등급 배지색 텍스트
        cell_text(gt.cell(ri,2),e["gs"],9.5,True,GRADE_COLOR.get(e["grade"],DGRAY),bg=bg)
        cell_text(gt.cell(ri,3),str(w),14,True,NAVY,bg=bg)
        cell_text(gt.cell(ri,4),rate,11,True if dep else False, DGRAY if not dep else NAVY, bg=bg)
        for i in range(4):
            if rets is None:
                cell_text(gt.cell(ri,5+i),"–",10,False,LGRAY,bg=bg)
            else:
                cell_text(gt.cell(ri,5+i),rets[i],10,True,ret_color(rets[i]),bg=bg)
    ri=2
    for key,w in rows:
        fillrow(ri,"ETF",ETF[key],w,"–",RETURNS[key]); ri+=1
    fillrow(ri,"예금",DEPOSIT,100-p["risk"],DEPOSIT["rate"],None,dep=True)

    # 각주
    add_text(s,IN(0.5),IN(6.75),IN(7.7),IN(0.55),
        [("※ ETF 기간별 수익률은 고객 직접 확인 기준 · 기준일 별도 표기 · 과거 수익률이 미래 수익을 보장하지 않습니다.",8.5,False,RGBColor(0x7a,0x85,0x95)),
         ("   예금 금리는 세전 연이율(확정)이며 가입시점 고시금리 적용.",8.5,False,RGBColor(0x7a,0x85,0x95))])

    # ----- 우측 도넛 차트 -----
    cd=CategoryChartData(); cd.categories=["위험자산(ETF)","안전자산(예금)"]
    cd.add_series("배분",(p["risk"],100-p["risk"]))
    gx=IN(9.0); gy=IN(1.55); gw=IN(3.9); gh=IN(2.75)
    gframe=s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT,gx,gy,gw,gh,cd)
    chart=gframe.chart; chart.has_title=False
    chart.has_legend=True; chart.legend.position=XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout=False; chart.legend.font.size=Pt(10); chart.legend.font.name=FONT
    plot=chart.plots[0]; plot.has_data_labels=False
    ser=plot.series[0]
    ser.points[0].format.fill.solid(); ser.points[0].format.fill.fore_color.rgb=ORANGE
    ser.points[1].format.fill.solid(); ser.points[1].format.fill.fore_color.rgb=SAFE
    # 중앙 라벨
    add_text(s,gx,IN(2.35),gw,IN(0.6),[(f"{p['risk']}%",24,True,NAVY)],align=PP_ALIGN.CENTER)

    # ----- 투자 포인트 -----
    box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,IN(8.95),IN(4.5),IN(3.9),IN(1.95))
    box.fill.solid();box.fill.fore_color.rgb=WHITE; box.line.color.rgb=RGBColor(0xE4,0xE9,0xF1);box.line.width=Pt(1)
    box.shadow.inherit=False
    tb=s.shapes.add_textbox(IN(9.15),IN(4.62),IN(3.55),IN(1.75)); tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=0;tf.margin_top=0
    p0=tf.paragraphs[0]; r=p0.add_run(); r.text="투자 포인트"; style_run(r,14,True,NAVY)
    pts=[f"위험자산 {p['risk']}% / 안전자산 {100-p['risk']}% — {p['profile']}",
         "반도체·AI 성장 테마 집중, 예금으로 변동성 하단 방어",
         "예금은 원리금보장·예금자보호로 안전판 역할"]
    for t in pts:
        pp=tf.add_paragraph(); pp.space_before=Pt(4)
        r=pp.add_run(); r.text="• "+t; style_run(r,10.5,False,RGBColor(0x37,0x42,0x50))

    # 하단 바
    foot=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,IN(7.05),IN(13.333),IN(0.45)); fill(foot,RGBColor(0xF4,0xF6,0xFA))
    add_text(s,IN(0.5),IN(7.1),IN(5),IN(0.35),[("미래에셋증권 마포WM",11,True,NAVY)],anchor=MSO_ANCHOR.MIDDLE)
    add_text(s,IN(8.3),IN(7.1),IN(4.5),IN(0.35),[("퇴직연금 DC 제안서 · 2026년 7월",11,False,RGBColor(0x7a,0x85,0x95))],align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE)

for p in PORTS: content(p)
prs.save(OUT)
print("saved",OUT)
